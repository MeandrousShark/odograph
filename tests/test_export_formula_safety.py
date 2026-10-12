"""XLSX text is stored as text; CSV stays literal."""
from __future__ import annotations

import csv
import io
from datetime import date, datetime, timezone
from decimal import Decimal
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from app.expenses import build_expense_report
from app.export import HEADERS, to_csv, to_range_report_xlsx, to_report_xlsx, to_xlsx
from app.odometer import VehicleCoverage
from app.rates import YearRate
from app.report import build_annual_report, build_range_report
from test_streamed_exports import make_renderer, send_row, trip as streamed_trip

pytestmark = pytest.mark.unit

TZ = ZoneInfo("America/Los_Angeles")
RATES = {2026: YearRate(0.7250)}
PAYLOADS = ('=HYPERLINK("http://x","y")', "=1+1", "=cmd|' /C calc'!A0")


def _trip(i: int, payload: str) -> dict:
    return dict(
        started_at=datetime(2026, 6, 15, 15, 0, tzinfo=timezone.utc),
        ended_at=datetime(2026, 6, 15, 15, 30, tzinfo=timezone.utc),
        distance_m=1609.344, display_distance_m=1609.344, category="business",
        purpose=payload, notes=payload, has_gap=False, source="detected",
        start_place_name=payload, start_lat=47.6, start_lon=-122.3,
        end_place_name=payload, end_lat=47.7, end_lon=-122.2,
        vehicle_id=i, vehicle_name=payload,
    )


TRIPS = [_trip(i, payload) for i, payload in enumerate(PAYLOADS, 1)]


def _book(raw: bytes):
    return load_workbook(io.BytesIO(raw))


def _assert_no_formulas(wb) -> None:
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (ws.title, cell.coordinate, cell.value)
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    assert cell.data_type == "s", (ws.title, cell.coordinate)


def _assert_trips_sheet(ws) -> None:
    col = {name: i for i, name in enumerate(HEADERS)}
    for row_index, payload in enumerate(PAYLOADS, 2):
        row = ws[row_index]
        for name in ("Start location", "End location", "Vehicle", "Purpose", "Notes"):
            cell = row[col[name]]
            assert (cell.value, cell.data_type) == (payload, "s"), (name, cell.coordinate)
        assert row[col["Distance (mi)"]].data_type == "n"
        assert row[col["Distance (km)"]].data_type == "n"
        assert row[col["Deduction ($)"]].data_type == "n"
        assert row[col["Deduction ($)"]].number_format == '"$"#,##0.00'
        assert row[col["Date"]].value == "2026-06-15"


def _assert_summary_text(ws) -> None:
    values = {c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
    assert set(PAYLOADS) <= values
    for row in ws.iter_rows():
        for cell in row:
            if cell.value in PAYLOADS:
                assert cell.data_type == "s"


def test_trips_xlsx_stores_formula_like_text_as_text():
    wb = _book(to_xlsx(TRIPS, RATES, TZ))
    _assert_no_formulas(wb)
    _assert_trips_sheet(wb["Trips"])


def _coverage():
    return [
        VehicleCoverage(
            vehicle_name=payload, span_start=datetime(2026, 3, 1, tzinfo=TZ),
            span_end=datetime(2026, 9, 1, tzinfo=TZ), coverage=0.8,
            gap_m=20 * 1609.344, fully_bracketed=False,
        )
        for payload in PAYLOADS
    ]


def test_annual_report_workbook_stores_formula_like_text_as_text():
    expenses = [{
        "vehicle_id": i, "vehicle_name": payload, "incurred_on": date(2026, 6, 1),
        "category": "fuel", "amount": Decimal("100.00"),
        "treatment": "business_use_allocated", "notes": payload,
    } for i, payload in enumerate(PAYLOADS, 1)]
    report = build_annual_report(TRIPS, RATES, TZ, 2026)
    expense_report = build_expense_report(2026, TRIPS, expenses, [], RATES, TZ)
    wb = _book(to_report_xlsx(report, TRIPS, RATES, TZ, _coverage(), expense_report, expenses))

    _assert_no_formulas(wb)
    _assert_summary_text(wb["Summary"])
    coverage_cells = [
        c for row in wb["Summary"].iter_rows() for c in row
        if isinstance(c.value, str) and "GPS captured" in c.value
    ]
    assert len(coverage_cells) == len(PAYLOADS)
    for cell, payload in zip(coverage_cells, PAYLOADS):
        assert cell.value.startswith(payload + ": GPS captured")
        assert cell.data_type == "s"
    _assert_trips_sheet(wb["Trips"])
    ws = wb["Expenses"]
    for row_index, payload in enumerate(PAYLOADS, 2):
        row = ws[row_index]
        assert (row[1].value, row[1].data_type) == (payload, "s")
        assert (row[5].value, row[5].data_type) == (payload, "s")
        assert row[0].data_type == "d"
        assert row[3].data_type == "n"
        assert row[3].number_format == '"$"#,##0.00'


def test_range_report_workbook_stores_formula_like_text_as_text():
    report = build_range_report(TRIPS, RATES, TZ, date(2026, 4, 1), date(2026, 6, 30))
    wb = _book(to_range_report_xlsx(report, TRIPS, RATES, TZ))

    _assert_no_formulas(wb)
    _assert_summary_text(wb["Summary"])
    _assert_trips_sheet(wb["Trips"])


def test_generic_streamed_export_stores_formula_like_text_as_text(tmp_path):
    renderer, reservation = make_renderer(tmp_path, "xlsx")
    try:
        for i, payload in enumerate(PAYLOADS):
            row = streamed_trip(i)
            row.update(
                vehicle_name=payload, purpose=payload, notes=payload,
                start_place_name=payload, end_place_name=payload,
                start_address=None, display_distance_m=1609.344,
                category="business", exclusion=None,
            )
            send_row(renderer, row)
        raw = renderer.budget.path(renderer.render()["path"]).read_bytes()
    finally:
        renderer.close()
        renderer.budget.close()
        reservation.release()
    wb = _book(raw)
    _assert_no_formulas(wb)
    ws = wb["Trips"]
    col = {name: i for i, name in enumerate(HEADERS)}
    for row_index, payload in enumerate(PAYLOADS, 2):
        row = ws[row_index]
        for name in ("Start location", "End location", "Vehicle", "Purpose", "Notes"):
            assert (row[col[name]].value, row[col[name]].data_type) == (payload, "s")
        assert row[col["Distance (mi)"]].data_type == "n"
        assert row[col["Deduction ($)"]].data_type == "n"


def test_csv_writes_formula_like_text_literally():
    rows = list(csv.reader(io.StringIO(to_csv(TRIPS, RATES, TZ).decode())))
    col = {name: i for i, name in enumerate(HEADERS)}
    for row, payload in zip(rows[1:], PAYLOADS):
        for name in ("Start location", "End location", "Vehicle", "Purpose", "Notes"):
            assert row[col[name]] == payload
    assert "=1+1" in to_csv(TRIPS, RATES, TZ).decode()
