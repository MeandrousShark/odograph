"""CSV/XLSX export and the annual/range report workbooks.
`build_export_rows` is pure — no I/O, no DB/openpyxl imports — so it's
unit-testable without a workbook or a filesystem. `to_csv`/`to_xlsx` are thin,
in-memory writers around it; `to_report_xlsx` adds a Summary sheet ahead
of the same Trips sheet, so the report carries its own audit-appendix detail.
`to_range_report_xlsx` reuses the same sheet writers for a range report,
without the annual-only Expenses/odometer sections.
"""
from __future__ import annotations

import csv
import io
from decimal import Decimal
from zoneinfo import ZoneInfo

from app.expenses import (
    CATEGORY_LABELS,
    TREATMENT_LABELS,
    ExpenseReport,
    comparison_caveat_lines,
    comparison_status,
)
from app.odometer import VehicleCoverage, coverage_line
from app.places_desc import describe_endpoint
from app.rates import METERS_PER_MILE, YearRate, deduction
from app.report import MONTH_ABBR, AnnualReport, RangeReport, caveat_lines, format_rate_periods, range_label

HEADERS = (
    "Date", "Start", "End", "Duration", "Start location", "End location",
    "Distance (mi)", "Distance (km)", "Category", "Vehicle", "Purpose", "Notes", "Gap", "Source",
    "Deduction ($)",
)


def _duration_str(started_at, ended_at) -> str:
    secs = int((ended_at - started_at).total_seconds())
    h, m = divmod(secs // 60, 60)
    return f"{h}h {m:02d}m" if h else f"{m}m"


def build_export_rows(trips: list[dict], rates: dict[int, YearRate], tz: ZoneInfo) -> list[list]:
    """One row per trip, in `HEADERS` order. `trips` rows are expected to
    carry the same keys `TRIP_COLUMNS` selects (including start/end place
    names). Uses `display_distance_m` — snapped distance when available,
    raw `distance_m` as fallback — so exports match what's shown on the
    trip list.
    """
    rows = []
    for t in trips:
        local_start = t["started_at"].astimezone(tz)
        local_end = t["ended_at"].astimezone(tz)
        distance_m = t["display_distance_m"]
        is_business = t["category"] == "business"
        ded = (
            deduction(distance_m, local_start.year, rates, local_start.month)
            if is_business else None
        )
        rows.append([
            local_start.strftime("%Y-%m-%d"),
            local_start.strftime("%H:%M"),
            local_end.strftime("%H:%M"),
            _duration_str(t["started_at"], t["ended_at"]),
            describe_endpoint(
                t.get("start_place_name"), t.get("start_lat"), t.get("start_lon"),
                t.get("start_address"),
            ),
            describe_endpoint(
                t.get("end_place_name"), t.get("end_lat"), t.get("end_lon"),
                t.get("end_address"),
            ),
            round(distance_m / METERS_PER_MILE, 1),
            round(distance_m / 1000.0, 1),
            t["category"],
            t.get("vehicle_name") or "",
            t.get("purpose") or "",
            t.get("notes") or "",
            "yes" if t.get("has_gap") else "",
            t["source"],
            round(ded, 2) if ded is not None else "",
        ])
    return rows


def to_csv(trips: list[dict], rates: dict[int, YearRate], tz: ZoneInfo) -> bytes:
    rows = build_export_rows(trips, rates, tz)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(HEADERS)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


def _populate_trips_sheet(ws, trips: list[dict], rates: dict[int, YearRate], tz: ZoneInfo) -> None:
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    rows = build_export_rows(trips, rates, tz)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)

    mi_col, km_col, ded_col = 7, 8, 15
    total_mi = sum(r[mi_col - 1] for r in rows)
    total_km = sum(r[km_col - 1] for r in rows)
    total_ded = sum(r[ded_col - 1] for r in rows if isinstance(r[ded_col - 1], (int, float)))
    totals = [""] * len(HEADERS)
    totals[0] = "Total"
    totals[mi_col - 1] = round(total_mi, 1)
    totals[km_col - 1] = round(total_km, 1)
    totals[ded_col - 1] = round(total_ded, 2)
    ws.append(totals)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    ded_letter = get_column_letter(ded_col)
    for row_idx in range(2, ws.max_row + 1):
        cell = ws[f"{ded_letter}{row_idx}"]
        if isinstance(cell.value, (int, float)):
            cell.number_format = '"$"#,##0.00'


def to_xlsx(trips: list[dict], rates: dict[int, YearRate], tz: ZoneInfo) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Trips"
    _populate_trips_sheet(ws, trips, rates, tz)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_summary_sheet(
    ws, report: AnnualReport, odometer_coverage: list[VehicleCoverage] | None = None,
    expense_report: ExpenseReport | None = None, title: str | None = None,
) -> None:
    """The report's headline numbers, laid out for a quick read rather than
    as a data table — a distinct shape from the Trips sheet's per-row detail.
    `title` defaults to the annual report's own heading; `to_range_report_xlsx`
    passes a `range_label`-derived one instead, so the two exports share this
    writer without the range report's Summary sheet ever calling itself annual.
    """
    from openpyxl.styles import Font

    ws.append([title or f"Annual Mileage Report — {report.year}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Business miles", round(report.business_m / METERS_PER_MILE, 1)])
    ws.append(["Personal miles", round(report.personal_m / METERS_PER_MILE, 1)])
    ws.append(["Total miles", round(report.total_m / METERS_PER_MILE, 1)])
    ws.append([
        "Business share",
        f"{report.business_pct:.1f}%" if report.business_pct is not None else "—",
    ])
    ws.append(["Rate(s) applied", format_rate_periods(report.rate_periods)])
    ws.append([
        "Total deduction",
        f"${report.total_deduction:,.2f}" if report.total_deduction is not None else "—",
    ])
    ws.append(["Trip count", report.trip_count])
    ws.append([])

    ws.append(["Month", "Trips", "Business (mi)", "Rate ($/mi)", "Deduction ($)"])
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    for month in report.months:
        ws.append([
            MONTH_ABBR[month.month],
            month.trip_count,
            round(month.business_m / METERS_PER_MILE, 1),
            round(month.rate_per_mi, 4) if month.rate_per_mi is not None else "—",
            round(month.deduction, 2) if month.deduction is not None else "—",
        ])

    if report.caveats.any:
        ws.append([])
        ws.append(["Caveats — review before filing"])
        ws[f"A{ws.max_row}"].font = Font(bold=True)
        for line in caveat_lines(report.caveats, report.year):
            ws.append([line])

    # Appended after the caveats block (rather than, say, right after the
    # month table) so the summary sheet's fixed row positions (title/headline
    # rows, month-table header/rows) stay put regardless of how many vehicles
    # a year has.
    if report.by_vehicle:
        ws.append([])
        ws.append(["By vehicle", "Business (mi)", "Total (mi)", "Deduction ($)"])
        header_row = ws.max_row
        for cell in ws[header_row]:
            cell.font = Font(bold=True)
        for v in report.by_vehicle:
            ws.append([
                v.vehicle_name,
                round(v.business_m / METERS_PER_MILE, 1),
                round(v.total_m / METERS_PER_MILE, 1),
                round(v.deduction, 2) if v.deduction is not None else "—",
            ])

    # Appended after "By vehicle" for the same fixed-row-position reason;
    # one text line per vehicle via `coverage_line` (not a data table) so
    # this can't drift in wording from the HTML report's own use of it.
    if odometer_coverage:
        ws.append([])
        ws.append(["Odometer reconciliation"])
        ws[f"A{ws.max_row}"].font = Font(bold=True)
        for line in odometer_coverage:
            ws.append([coverage_line(line)])

    if expense_report and expense_report.comparisons:
        ws.append([])
        ws.append(["Standard vs. actual expense estimate"])
        ws[f"A{ws.max_row}"].font = Font(bold=True)
        ws.append([
            "Vehicle", "Business (mi)", "Total (mi)", "Business use", "Total-mile source",
            "Shared expenses ($)", "Fully business ($)", "Standard estimate ($)",
            "Actual estimate ($)", "Larger estimate", "Status",
        ])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        for line in expense_report.comparisons:
            if line.larger_estimate == "tie":
                larger = "Tie"
            elif line.larger_estimate:
                larger = f"{line.larger_estimate.title()} by ${line.difference:,.2f}"
                if line.provisional:
                    larger += " (provisional)"
            else:
                larger = "—"
            ws.append([
                line.vehicle_name,
                round(line.business_m / METERS_PER_MILE, 1),
                round(line.denominator_m / METERS_PER_MILE, 1),
                f"{line.business_pct * 100:.1f}%" if line.business_pct is not None else "—",
                "Odometer" if line.denominator_source == "odometer" else "GPS detected",
                float(line.allocated_expenses),
                float(line.fully_business_expenses),
                float(line.standard_total) if line.standard_total is not None else "—",
                float(line.actual_total) if line.actual_total is not None else "—",
                larger,
                comparison_status(line),
            ])
        ws.append([])
        ws.append(["Actual-expense caveats — review before filing"])
        ws[f"A{ws.max_row}"].font = Font(bold=True)
        for line in comparison_caveat_lines(expense_report.comparisons):
            ws.append([line])
        ws.append(["IRS guidance: Publication 463 and Publication 946."])

    for col, width in zip("ABCDEFGHIJK", (32, 18, 14, 16, 18, 19, 18, 20, 18, 28, 18)):
        ws.column_dimensions[col].width = width


def _populate_expenses_sheet(ws, expenses: list[dict]) -> None:
    from openpyxl.styles import Font

    headers = ("Date", "Vehicle", "Category", "Amount ($)", "Tax treatment", "Notes")
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    total = Decimal("0")
    for expense in expenses:
        amount = Decimal(str(expense["amount"]))
        total += amount
        ws.append([
            expense["incurred_on"],
            expense["vehicle_name"],
            CATEGORY_LABELS[expense["category"]],
            float(amount),
            TREATMENT_LABELS[expense["treatment"]],
            expense.get("notes") or "",
        ])
        ws.cell(ws.max_row, 4).number_format = '"$"#,##0.00'
    ws.append(["Total", "", "", float(total.quantize(Decimal("0.01"))), "", ""])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.cell(ws.max_row, 4).number_format = '"$"#,##0.00'
    for col, width in zip("ABCDEF", (14, 22, 24, 14, 24, 36)):
        ws.column_dimensions[col].width = width


def to_report_xlsx(
    report: AnnualReport, trips: list[dict], rates: dict[int, YearRate], tz: ZoneInfo,
    odometer_coverage: list[VehicleCoverage] | None = None,
    expense_report: ExpenseReport | None = None,
    expenses: list[dict] | None = None,
) -> bytes:
    """Summary sheet with the headline numbers/month table/caveats, plus a
    Trips sheet (the same detail `to_xlsx` writes, filtered to the report's
    year) as an audit appendix backing the summary. When an expense report is
    supplied, an Expenses sheet carries the ledger behind the comparison.
    `odometer_coverage` is computed by the caller (app.ui's report routes),
    same as `report` itself — kept a plain parameter here rather than folded
    into `AnnualReport` so `build_annual_report`'s signature stays untouched.
    """
    from openpyxl import Workbook

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"
    _write_summary_sheet(summary_ws, report, odometer_coverage, expense_report)

    trips_ws = wb.create_sheet("Trips")
    _populate_trips_sheet(trips_ws, trips, rates, tz)

    if expense_report is not None:
        expenses_ws = wb.create_sheet("Expenses")
        _populate_expenses_sheet(expenses_ws, expenses or [])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_range_report_xlsx(
    report: RangeReport, trips: list[dict], rates: dict[int, YearRate], tz: ZoneInfo
) -> bytes:
    """Range-report analog of `to_report_xlsx` — reuses the same
    `_write_summary_sheet`/`_populate_trips_sheet` writers so the two exports
    can't drift on layout, with no `odometer_coverage`/`expense_report`
    arguments to pass through (the range report deliberately carries no
    odometer-coverage or standard-vs-actual section, so there's nothing
    app/ui.py's range routes need to fetch for those sheets).
    `trips` is filtered here to those whose local start date falls in
    `[report.start, report.end]` — the same rule `build_range_report` used to
    fold the summary numbers above — so every Trips-sheet row backs a row
    already counted in the summary, even though the caller's DB query only
    coarsely pre-filters (see `app/ui.py`'s `_fetch_range_trips`).
    """
    from openpyxl import Workbook

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"
    _write_summary_sheet(
        summary_ws, report, title=f"Mileage Report — {range_label(report.start, report.end)}"
    )

    trips_ws = wb.create_sheet("Trips")
    in_range_trips = [
        t for t in trips if report.start <= t["started_at"].astimezone(tz).date() <= report.end
    ]
    _populate_trips_sheet(trips_ws, in_range_trips, rates, tz)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
