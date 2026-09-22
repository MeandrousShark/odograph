"""DB-backed tests for the arbitrary date-range/quarterly report routes:
`/report/range` and `/report/range/export`. Route
handlers are invoked directly, same convention `tests/test_odometer_db.py`/
`tests/test_review_db.py` use. Also covers `/report/{year}`'s next-year
guard's timezone wiring (`test_report_page_...`), since this is the DB-backed
report-route file with the `_endpoint`/`_request` harness that test needs.
"""
from __future__ import annotations

import asyncio
import os
from datetime import datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.db import make_pool
from app.account_context import account_id
from personal_support import personal_request
from app.export import HEADERS
from app.main import make_templates
from app.ui import make_router
import app.ui.reports as ui
from conftest import reset_account_db

TEST_DB = os.environ.get("TEST_DATABASE_URL")
pytestmark = pytest.mark.skipif(not TEST_DB, reason="set TEST_DATABASE_URL to run DB-backed tests")

TZ = ZoneInfo("America/Los_Angeles")
USER = {"sub": "test"}


def _endpoint(path: str):
    for route in make_router().routes:
        if getattr(route, "path", None) == path:
            return route.endpoint
    raise AssertionError(f"route {path} missing")


def _request(pool):
    config = SimpleNamespace(display_tz=TZ, app_version="test")
    return personal_request(SimpleNamespace(
        app=SimpleNamespace(state=SimpleNamespace(
            pool=pool, config=config, templates=make_templates(config),
        )),
        session={"csrf": "token"},
    ))


async def _create_vehicle(conn, name: str) -> int:
    cur = await conn.execute('INSERT INTO vehicles (account_id, name) VALUES (%s, %s) RETURNING id', (
                                                                                                         account_id(conn),
                                                                                                         name,
                                                                                                     ))
    return (await cur.fetchone())[0]


async def _insert_trip(
    conn, vehicle_id: int, started_at: datetime, distance_m: float, category: str = "business",
) -> None:
    await conn.execute(
        "INSERT INTO trips (account_id, device, source, started_at, ended_at, distance_m, "
        "vehicle_id, category) VALUES (%s, 'manual', 'manual', %s, %s, %s, %s, %s)",
        (
            account_id(conn),
            started_at,
            started_at + timedelta(minutes=15),
            distance_m,
            vehicle_id,
            category,
        ),
    )


async def _range_page_scenario():
    raw_pool = make_pool(TEST_DB)
    await raw_pool.open(wait=True)
    try:
        pool = await reset_account_db(raw_pool)
        async with pool.connection() as conn:
            truck_id = await _create_vehicle(conn, "Truck")
            # March 31 and July 1 are outside the Apr 1 - Jun 30 range under
            # test and must not contribute to its totals; April 15/June 30
            # are the in-range trips whose miles/deduction the page must show.
            await _insert_trip(
                conn, truck_id, datetime(2026, 3, 31, 20, tzinfo=timezone.utc), 10 * 1609.344
            )
            await _insert_trip(
                conn, truck_id, datetime(2026, 4, 15, 19, tzinfo=timezone.utc), 20 * 1609.344
            )
            await _insert_trip(
                conn, truck_id, datetime(2026, 6, 30, 19, tzinfo=timezone.utc), 30 * 1609.344
            )
            await _insert_trip(
                conn, truck_id, datetime(2026, 7, 1, 19, tzinfo=timezone.utc), 40 * 1609.344
            )

        request = _request(pool)
        page = _endpoint("/report/range")
        response = await page(request, from_="2026-04-01", to="2026-06-30", user=USER)
        assert response.status_code == 200
        body = response.body.decode()
        assert "2026 Q2" in body
        # 50 mi total business (20 + 30) at 2026's seeded $0.7250/mi rate.
        assert "50.0 mi" in body
        assert "$36.25" in body

        # Cross-check against the pure fold directly (not just page text) so
        # a future markup change can't silently mask a totals mismatch.
        from datetime import date

        from psycopg.rows import dict_row

        from app.rates import load_rates
        from app.report import build_range_report

        async with pool.connection() as conn:
            cur = conn.cursor(row_factory=dict_row)
            await cur.execute(
                "SELECT id, device, source::text AS source, started_at, ended_at, "
                "distance_m AS display_distance_m, category::text AS category, "
                "vehicle_id, false AS has_gap, 'ok'::text AS snap_status "
                "FROM trips ORDER BY started_at"
            )
            trips = await cur.fetchall()
            rates = await load_rates(conn)
        expected = build_range_report(trips, rates, TZ, date(2026, 4, 1), date(2026, 6, 30))
        assert expected.business_m == pytest.approx(50 * 1609.344)
        assert expected.total_deduction == pytest.approx(36.25)
    finally:
        await raw_pool.close()


def test_range_report_page_200_matches_pure_fold_totals():
    asyncio.run(_range_page_scenario())


async def _range_page_invalid_params_scenario():
    raw_pool = make_pool(TEST_DB)
    await raw_pool.open(wait=True)
    try:
        pool = await reset_account_db(raw_pool)
        request = _request(pool)
        page = _endpoint("/report/range")

        with pytest.raises(HTTPException) as malformed:
            await page(request, from_="not-a-date", to="2026-06-30", user=USER)
        assert malformed.value.status_code == 400

        with pytest.raises(HTTPException) as missing:
            await page(request, from_="", to="2026-06-30", user=USER)
        assert missing.value.status_code == 400

        with pytest.raises(HTTPException) as reversed_range:
            await page(request, from_="2026-06-30", to="2026-04-01", user=USER)
        assert reversed_range.value.status_code == 400

        with pytest.raises(HTTPException) as cross_year:
            await page(request, from_="2025-12-15", to="2026-01-15", user=USER)
        assert cross_year.value.status_code == 400
    finally:
        await raw_pool.close()


def test_range_report_page_400_for_each_invalid_param_case():
    asyncio.run(_range_page_invalid_params_scenario())


async def _range_export_invalid_params_scenario():
    raw_pool = make_pool(TEST_DB)
    await raw_pool.open(wait=True)
    try:
        pool = await reset_account_db(raw_pool)
        request = _request(pool)
        export = _endpoint("/report/range/export")

        with pytest.raises(HTTPException) as malformed:
            await export(request, from_="2026-13-40", to="2026-06-30", user=USER)
        assert malformed.value.status_code == 400

        with pytest.raises(HTTPException) as reversed_range:
            await export(request, from_="2026-06-30", to="2026-04-01", user=USER)
        assert reversed_range.value.status_code == 400

        with pytest.raises(HTTPException) as cross_year:
            await export(request, from_="2025-12-15", to="2026-01-15", user=USER)
        assert cross_year.value.status_code == 400
    finally:
        await raw_pool.close()


def test_range_report_export_400_for_each_invalid_param_case():
    asyncio.run(_range_export_invalid_params_scenario())


async def _range_export_scenario():
    raw_pool = make_pool(TEST_DB)
    await raw_pool.open(wait=True)
    try:
        pool = await reset_account_db(raw_pool)
        async with pool.connection() as conn:
            truck_id = await _create_vehicle(conn, "Truck")
            await _insert_trip(
                conn, truck_id, datetime(2026, 4, 15, 19, tzinfo=timezone.utc), 20 * 1609.344,
            )
            await _insert_trip(
                conn, truck_id, datetime(2026, 7, 1, 19, tzinfo=timezone.utc), 40 * 1609.344,
            )

        request = _request(pool)
        export = _endpoint("/report/range/export")
        response = await export(request, from_="2026-04-01", to="2026-06-30", user=USER)
        assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert response.headers["Content-Disposition"] == 'attachment; filename="mileage-report-2026-Q2.xlsx"'

        wb = load_workbook(BytesIO(response.body))
        assert wb.sheetnames == ["Summary", "Trips"]
        assert wb["Summary"]["A1"].value == "Mileage Report: 2026 Q2"

        trips_ws = wb["Trips"]
        dates = [
            trips_ws.cell(row, 1).value for row in range(2, trips_ws.max_row)  # exclude totals row
        ]
        assert dates == ["2026-04-15"]  # the July trip is filtered out of the Trips sheet
    finally:
        await raw_pool.close()


def test_range_report_export_media_type_filename_and_row_filtering():
    asyncio.run(_range_export_scenario())


async def _range_export_endpoint_label_scenario():
    raw_pool = make_pool(TEST_DB)
    await raw_pool.open(wait=True)
    try:
        pool = await reset_account_db(raw_pool)
        async with pool.connection() as conn:
            truck_id = await _create_vehicle(conn, "Truck")
            # A route_mode=none manual trip: source manual, no place id or
            # geometry on either endpoint, so migrations/025_manual_trip_labels.sql's
            # constraints allow both labels. This proves TRIP_COLUMNS'
            # COALESCE(start_label, saved-place name) actually resolves into
            # start_place_name/end_place_name for a real row selected by the
            # real query, not just a hand-built trip dict.
            await conn.execute(
                "INSERT INTO trips (account_id, device, source, started_at, ended_at, distance_m, "
                "vehicle_id, category, start_label, end_label) VALUES (%s, 'manual', 'manual', %s, "
                "%s, %s, %s, %s, %s, %s)",
                (
                    account_id(conn),
                    datetime(2026, 4, 15, 19, tzinfo=timezone.utc),
                    datetime(2026, 4, 15, 19, 15, tzinfo=timezone.utc),
                    20 * 1609.344,
                    truck_id,
                    "business",
                    "Grandma's house",
                    "Trailhead parking",
                ),
            )

        request = _request(pool)
        export = _endpoint("/report/range/export")
        response = await export(request, from_="2026-04-01", to="2026-06-30", user=USER)

        wb = load_workbook(BytesIO(response.body))
        trips_ws = wb["Trips"]
        assert trips_ws.cell(2, HEADERS.index("Start location") + 1).value == "Grandma's house"
        assert trips_ws.cell(2, HEADERS.index("End location") + 1).value == "Trailhead parking"
    finally:
        await raw_pool.close()


def test_range_report_export_resolves_endpoint_labels_through_trip_columns():
    asyncio.run(_range_export_endpoint_label_scenario())


class _FrozenDatetime(datetime):
    """`datetime.now(tz)` fixed at 2027-01-01 03:00 UTC (2026-12-31 19:00 in
    `America/Los_Angeles`) -- the instant a UTC-based "current year" would
    already read 2027 while the display timezone's hasn't rolled over yet.
    Subclasses the real `datetime`, not a bare stub, so every other
    `datetime(...)` construction `app.ui` does elsewhere (e.g. `_fetch_range_trips`'s
    day boundaries) keeps working unchanged; only `.now()` is fixed.
    """

    _instant = datetime(2027, 1, 1, 3, 0, tzinfo=timezone.utc)

    @classmethod
    def now(cls, tz=None):
        return cls._instant if tz is None else cls._instant.astimezone(tz)


async def _report_page_year_boundary_scenario():
    raw_pool = make_pool(TEST_DB)
    await raw_pool.open(wait=True)
    try:
        pool = await reset_account_db(raw_pool)
        request = _request(pool)
        page = _endpoint("/report/{year}")

        # The operator's local year is still 2026 at this instant, so 2027
        # hasn't started for them -- the control offering it must be
        # disabled even though UTC already reads January 2027.
        current = await page(request, year=2026, user=USER)
        current_body = current.body.decode()
        assert '<span class="report-year-next" aria-disabled="true">2027 →</span>' in current_body
        assert 'href="/report/2027"' not in current_body
        assert '<a href="/report/2025">← 2025</a>' in current_body

        # 2026 itself, one year forward from 2025, has already started (most
        # of it happened before this instant) -- its own control must stay
        # enabled, which is what catches an off-by-one the other way.
        past = await page(request, year=2025, user=USER)
        past_body = past.body.decode()
        assert '<a href="/report/2026">2026 →</a>' in past_body
        assert 'aria-disabled="true"' not in past_body
    finally:
        await raw_pool.close()


def test_report_page_next_year_guard_uses_display_timezone_at_new_years_eve_boundary(monkeypatch):
    # Proves app/ui/reports.py's report_page route passes a tz-localized `now` (not
    # a naive/UTC one) into next_year_disabled: patching ui.datetime.now to a
    # fixed instant and letting the route localize it via `datetime.now(tz)`
    # is the only way this test can distinguish the two -- a plain
    # `datetime.now()` call site would see the frozen instant's UTC year
    # (2027) instead and fail the assertions above.
    monkeypatch.setattr(ui, "datetime", _FrozenDatetime)
    asyncio.run(_report_page_year_boundary_scenario())
