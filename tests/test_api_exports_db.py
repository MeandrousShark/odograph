"""API v1 exports and report workbooks through the real application: the
same files as the web's downloads, bearer-only identity, the binding
re-check at the snapshot, and envelopes for every failure."""
from __future__ import annotations

import asyncio
import json
import os
import re
import zipfile
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from openpyxl import load_workbook

from app.account_context import AccountPool, AccountPrincipal, account_id, control_connection
from app.api_errors import ERRORS
from app.api_grants import create_grant
from app.capacity import AdmissionManager, CapacityBusy
from app.db import make_pool
from app.main import create_app
from app.preparation import PreparationOperation
from app.preparation_resources import PreparationBusy, PreparationResourceError
from app.report_preparation import AccountBindingChanged
from auth_db_fixtures import auth_config
from conftest import add_test_account, reset_db
from test_report_projection_db import TZ

TEST_DB = os.environ.get("TEST_DATABASE_URL")
pytestmark = [pytest.mark.capacity_contract,
              pytest.mark.skipif(not TEST_DB, reason="requires disposable PostGIS")]
FIXTURES = Path(__file__).parent / "fixtures" / "api_v1"
OTHER = 77
OTHER_MARKER = "FOREIGN API EXPORT MARKER"
CSV = "text/csv; charset=utf-8"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@asynccontextmanager
async def application(tmp_path):
    raw = make_pool(TEST_DB)
    await raw.open(wait=True)
    try:
        await reset_db(raw)
        config = auth_config(TEST_DB, dev_no_auth=False, initial_admin_signup=True,
                             preparation_spool_dir=str(tmp_path / "spool"), raw_message_retention_days=0)
        app = create_app(config)
        async with app.router.lifespan_context(app):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),
                                         base_url="https://test") as client:
                page = await client.get("/signup")
                csrf = re.search(r'name="csrf_token" value="([^"]+)"', page.text).group(1)
                response = await client.post("/signup", data={
                    "email": "api-export-owner@example.invalid", "password": "api export test password",
                    "password_confirm": "api export test password", "csrf_token": csrf,
                    "display_timezone": str(TZ)})
                assert response.status_code == 303
                async with raw.connection() as conn:
                    owner = (await (await conn.execute(
                        "SELECT id FROM accounts WHERE email = %s",
                        ("api-export-owner@example.invalid",))).fetchone())[0]
                pool = AccountPool(app.state.runtime_pool, AccountPrincipal(owner, True, 1))
                async with pool.connection() as conn:
                    await _seed(conn, "Route car 車", "Client visit")
                other = await add_test_account(raw, OTHER)
                async with other.connection() as conn:
                    await conn.execute("UPDATE account_settings SET display_tz = %s WHERE account_id = %s",
                                       (str(TZ), OTHER))
                    await _seed(conn, "Their car", OTHER_MARKER)
                tokens = {}
                async with control_connection(app.state.control_pool) as conn:
                    for account in (owner, OTHER):
                        tokens[account] = (await create_grant(conn, account, 1, "Phone")).access_token
                yield SimpleNamespace(raw=raw, app=app, client=client, owner=owner, tokens=tokens,
                                      config=config)
        snapshot = app.state.capacity.snapshot()
        assert not snapshot["leases"]
        assert all(not lane["active"] and not lane["pending"]
                   for lane in snapshot.values() if isinstance(lane, dict))
    finally:
        await raw.close()


async def _seed(conn, vehicle_name, purpose):
    owner = account_id(conn)
    vehicle = (await (await conn.execute(
        "SELECT id FROM vehicles WHERE account_id = %s AND is_default", (owner,))).fetchone())[0]
    await conn.execute("UPDATE vehicles SET name = %s WHERE id = %s", (vehicle_name, vehicle))
    for index, category in enumerate(("business", "personal", "business", "unclassified")):
        start = datetime(2026, 4, 15 + index, 9, tzinfo=TZ)
        await conn.execute(
            "INSERT INTO trips (account_id, device, source, started_at, ended_at, distance_m, category, "
            "exclusion, vehicle_id, start_label, end_label, purpose, notes) "
            "VALUES (%s, 'route', 'manual', %s, %s, %s, %s, %s, %s, %s, 'End', %s, %s)",
            (owner, start, start + timedelta(minutes=30), 1609.344 * (index + 1), category,
             "not_deductible" if index == 2 else None, vehicle if index != 3 else None,
             "Start 車", purpose, "車,%\"\r\n" * 50 + ("literal%_\\" if index % 2 else "other")))
    await conn.execute(
        "INSERT INTO expenses (account_id, vehicle_id, incurred_on, category, amount, treatment, notes) "
        "VALUES (%s, %s, %s, 'parking', 12.34, 'fully_business', 'Garage 車')",
        (owner, vehicle, date(2026, 4, 15)))


def _bearer(state, account=None):
    return {"Authorization": f"Bearer {state.tokens[account or state.owner]}"}


def _envelope(code, fields=None):
    return {"error": {"code": code, "message": ERRORS[code][1], "fields": fields}}


def _fixture(name):
    return json.loads((FIXTURES / name).read_text())


def _xlsx_parts(raw):
    """Every zip member's bytes, in order. openpyxl stamps the wall-clock
    time into the zip entry headers and docProps/core.xml inside the
    isolated helper process, so only those two timestamps are dropped."""
    with zipfile.ZipFile(BytesIO(raw)) as book:
        parts = []
        for info in book.infolist():
            data = book.read(info)
            if info.filename == "docProps/core.xml":
                data, count = re.subn(rb"<dcterms:(created|modified)[^>]*>[^<]*</dcterms:\1>", b"", data)
                assert count == 2
            parts.append((info.filename, info.compress_type, data))
        return parts


def _cell_text(raw):
    """Every cell value in the workbook, as one string."""
    book = load_workbook(BytesIO(raw), read_only=True)
    try:
        return "\n".join(str(value) for sheet in book for row in sheet.iter_rows(values_only=True)
                         for value in row if value is not None)
    finally:
        book.close()


def _same_file(api, web, media_type, filename):
    assert web.status_code == 200 and api.status_code == 200
    if media_type == CSV:
        assert api.content == web.content
    else:
        assert _xlsx_parts(api.content) == _xlsx_parts(web.content)
    for response in (api, web):
        assert response.headers["content-type"] == media_type
        assert response.headers["content-length"] == str(len(response.content))
        assert response.headers["content-disposition"] == f'attachment; filename="{filename}"'
        assert response.headers["cache-control"] == "no-store, private"


def _observe(monkeypatch):
    operations = []
    original = PreparationOperation.start_helper

    async def start(operation, *args, **kwargs):
        operations.append(operation)
        return await original(operation, *args, **kwargs)

    monkeypatch.setattr(PreparationOperation, "start_helper", start)
    return operations


def _assert_clean(state, operations):
    snapshot = state.app.state.capacity.snapshot()
    assert snapshot["foreground"] == {"active": 0, "pending": 0}
    assert snapshot["leases"] == 0
    assert not list(Path(state.config.preparation_spool_dir).glob("op-*"))
    for operation in operations:
        assert operation.closed and not operation.directory.exists()
        assert operation.process.returncode is not None


# Each case shares one application across its requests: the database reset
# per case, not the request, is what these cost.
FILE_PATHS = [
    ("/api/v1/exports/trips", {"format": "csv"}, "/export"),
    ("/api/v1/reports/2026", {"format": "xlsx"}, "/report/2026/export"),
    ("/api/v1/reports/range", {"format": "xlsx", "from": "2026-04-01", "to": "2026-06-30"},
     "/report/range/export"),
]
WEB_FILE_PATHS = [
    ("/export", {"format": "csv"}),
    ("/report/2026/export", {}),
    ("/report/range/export", {"from": "2026-04-01", "to": "2026-06-30"}),
    ("/report/2026", {}),
]
EXPORT_FILTERS = [
    {},
    {"category": "business", "q": "literal%_\\"},
    {"from": "2026-04-16", "to": "2026-04-17", "exclusion": "none"},
    {"vehicle": "none"},
    {"exclusion": "not_deductible", "q": "  Client  "},
]
BAD_PARAMETERS = [
    ("/api/v1/exports/trips", {}, {"format": "This field is required."}),
    ("/api/v1/exports/trips", {"format": "pdf"}, {"format": "This value is not valid."}),
    ("/api/v1/exports/trips", {"format": "html"}, {"format": "This value is not valid."}),
    ("/api/v1/exports/trips", {"format": "csv", "category": "bogus"}, {"category": "This value is not valid."}),
    ("/api/v1/exports/trips", {"format": "csv", "exclusion": "all"}, {"exclusion": "This value is not valid."}),
    ("/api/v1/exports/trips", {"format": "csv", "from": "2026-4-1"}, {"from": "This value is not valid."}),
    ("/api/v1/exports/trips", {"format": "csv", "to": "2026-02-30"}, {"to": "This value is not valid."}),
    ("/api/v1/exports/trips", {"format": "csv", "to": "9999-12-31"}, {"to": "This value is not valid."}),
    ("/api/v1/exports/trips", {"format": "csv", "from": "1899-12-31"}, {"from": "This value is not valid."}),
    ("/api/v1/exports/trips", {"format": "csv", "vehicle": "car"}, {"vehicle": "This value is not valid."}),
    ("/api/v1/exports/trips", {"format": "csv", "vehicle": "0"}, {"vehicle": "This value is not valid."}),
    ("/api/v1/exports/trips", {"format": "csv", "q": "a\x00b"}, {"q": "This value is not valid."}),
    ("/api/v1/reports/2026", {}, {"format": "This field is required."}),
    ("/api/v1/reports/2026", {"format": "html"}, {"format": "This value is not valid."}),
    ("/api/v1/reports/2026", {"format": "csv"}, {"format": "This value is not valid."}),
    ("/api/v1/reports/1899", {"format": "xlsx"}, {"year": "This value is not valid."}),
    ("/api/v1/reports/9999", {"format": "xlsx"}, {"year": "This value is not valid."}),
    ("/api/v1/reports/twenty", {"format": "xlsx"}, {"year": "This value is not valid."}),
    ("/api/v1/reports/0002026", {"format": "xlsx"}, {"year": "This value is not valid."}),
    ("/api/v1/reports/02026", {"format": "xlsx"}, {"year": "This value is not valid."}),
    ("/api/v1/reports/+2026", {"format": "xlsx"}, {"year": "This value is not valid."}),
    ("/api/v1/reports/2026.0", {"format": "xlsx"}, {"year": "This value is not valid."}),
    ("/api/v1/reports/2_026", {"format": "xlsx"}, {"year": "This value is not valid."}),
    ("/api/v1/reports/%202026", {"format": "xlsx"}, {"year": "This value is not valid."}),
    ("/api/v1/reports/2026%20", {"format": "xlsx"}, {"year": "This value is not valid."}),
    ("/api/v1/reports/%D9%A2%D9%A0%D9%A2%D9%A6", {"format": "xlsx"}, {"year": "This value is not valid."}),
    ("/api/v1/reports/range", {"format": "xlsx"},
     {"from": "This field is required.", "to": "This field is required."}),
    ("/api/v1/reports/range", {"format": "html", "from": "2026-01-01", "to": "2026-03-31"},
     {"format": "This value is not valid."}),
    ("/api/v1/reports/range", {"format": "xlsx", "from": "2026-01-01", "to": "2026-3-31"},
     {"to": "This value is not valid."}),
    ("/api/v1/reports/range", {"format": "xlsx", "from": "2026-04-01", "to": "2027-03-31"},
     {"to": "Keep the range within one calendar year."}),
    ("/api/v1/reports/range", {"format": "xlsx", "from": "1899-01-01", "to": "1899-03-31"},
     {"from": "This value is not valid."}),
]


def test_trip_exports_match_the_web_files(tmp_path, monkeypatch):
    operations = _observe(monkeypatch)

    async def run():
        async with application(tmp_path) as state:
            async with state.raw.connection() as conn:
                vehicle = (await (await conn.execute(
                    "SELECT id FROM vehicles WHERE account_id = %s", (state.owner,))).fetchone())[0]
            for format in ("csv", "xlsx"):
                for filters in [*EXPORT_FILTERS, {"vehicle": str(vehicle)}]:
                    params = {"format": format, **filters}
                    web = await state.client.get("/export", params=params)
                    api = await state.client.get("/api/v1/exports/trips", params=params, headers=_bearer(state))
                    _same_file(api, web, CSV if format == "csv" else XLSX, f"trips.{format}")
                    assert OTHER_MARKER.encode() not in api.content
                    if format == "csv" and not filters:
                        assert api.content.count(b"Client visit") == 4
                    if format == "csv" and filters.get("vehicle") == str(vehicle):
                        assert api.content.count(b"Client visit") == 3
            assert len(operations) == 24
            _assert_clean(state, operations)
    asyncio.run(run())


def test_report_workbooks_match_the_web_files_and_stay_in_the_account(tmp_path, monkeypatch):
    operations = _observe(monkeypatch)

    async def run():
        async with application(tmp_path) as state:
            for api_path, web_path, params, filename in (
                ("/api/v1/reports/2026", "/report/2026/export", {}, "mileage-report-2026.xlsx"),
                ("/api/v1/reports/range", "/report/range/export", {"from": "2026-04-01", "to": "2026-06-30"},
                 "mileage-report-2026-Q2.xlsx"),
                ("/api/v1/reports/range", "/report/range/export", {"from": "2026-04-16", "to": "2026-04-17"},
                 "mileage-report-2026-04-16_2026-04-17.xlsx"),
            ):
                web = await state.client.get(web_path, params=params)
                api = await state.client.get(api_path, params={"format": "xlsx", **params}, headers=_bearer(state))
                _same_file(api, web, XLSX, filename)
                text = _cell_text(api.content)
                assert "Route car 車" in text and "Client visit" in text
                assert OTHER_MARKER not in text
            _assert_clean(state, operations)
            # Another account's token sees only that account's rows.
            theirs = await state.client.get("/api/v1/exports/trips", params={"format": "csv"},
                                            headers=_bearer(state, OTHER))
            assert theirs.status_code == 200
            assert theirs.content.count(OTHER_MARKER.encode()) == 4
            assert b"Client visit" not in theirs.content and "Route car 車".encode() not in theirs.content
            report = await state.client.get("/api/v1/reports/2026", params={"format": "xlsx"},
                                            headers=_bearer(state, OTHER))
            text = _cell_text(report.content)
            assert "Their car" in text and OTHER_MARKER in text
            assert "Route car 車" not in text and "Client visit" not in text
    asyncio.run(run())


def test_bad_parameters_missing_bearers_and_unknown_shapes_are_envelopes(tmp_path, monkeypatch):
    operations = _observe(monkeypatch)

    async def run():
        async with application(tmp_path) as state:
            for path, params, fields in BAD_PARAMETERS:
                response = await state.client.get(path, params=params, headers=_bearer(state))
                assert response.status_code == 422, (path, params)
                assert response.headers["content-type"] == "application/json"
                assert response.json() == _envelope("validation_failed", fields), (path, params)
            response = await state.client.get("/api/v1/exports/trips", params={"format": "pdf"},
                                              headers=_bearer(state))
            assert response.json() == _fixture("error_export_format_invalid.json")
            response = await state.client.get(
                "/api/v1/reports/range", params={"format": "xlsx", "from": "2026-05-01", "to": "2026-04-01"},
                headers=_bearer(state))
            assert response.status_code == 422
            assert response.json() == _fixture("error_report_range_reversed.json")
            # The client still carries the web session cookie from signup.
            assert state.client.cookies.get("session")
            for path, params, _ in FILE_PATHS:
                for headers in ({}, {"Authorization": "Bearer " + "A" * 43}, {"Authorization": "Basic eDp5"}):
                    response = await state.client.get(path, params=params, headers=headers)
                    assert response.status_code == 401
                    assert response.headers["www-authenticate"] == "Bearer"
                    assert response.json() == _fixture("error_invalid_token.json")
            # The HTML reports and the web's URL shapes are not API routes.
            for path in ("/api/v1/reports/2026/export", "/api/v1/reports/2026/html",
                         "/api/v1/reports/range/export", "/api/v1/report/2026"):
                response = await state.client.get(path, params={"format": "html"}, headers=_bearer(state))
                assert response.status_code == 404
                assert response.json() == _envelope("not_found")
            assert not operations
            _assert_clean(state, operations)
    asyncio.run(run())


def test_binding_change_after_admission_fails_at_the_snapshot(tmp_path, monkeypatch):
    operations = []
    original = PreparationOperation.start_helper

    async def run():
        async with application(tmp_path) as state:
            change = None

            async def start(operation, *args, **kwargs):
                operations.append(operation)
                session = await original(operation, *args, **kwargs)
                # Admission has passed; the snapshot's guard has not run yet.
                async with state.raw.connection() as conn:
                    await conn.execute(f"UPDATE accounts SET {change} WHERE id = %s", (state.owner,))
                return session
            monkeypatch.setattr(PreparationOperation, "start_helper", start)

            async def restore():
                async with state.raw.connection() as conn:
                    await conn.execute("UPDATE accounts SET auth_version = 1, is_enabled = true WHERE id = %s",
                                       (state.owner,))
            for change in ("auth_version = auth_version + 1", "is_enabled = false"):
                for path, params, _ in FILE_PATHS:
                    operations.clear()
                    response = await state.client.get(path, params=params, headers=_bearer(state))
                    assert response.status_code == 401
                    assert response.headers["www-authenticate"] == "Bearer"
                    assert response.json() == _envelope("invalid_token")
                    assert len(operations) == 1
                    _assert_clean(state, operations)
                    # The grant no longer authenticates either.
                    response = await state.client.get(path, params=params, headers=_bearer(state))
                    assert response.status_code == 401 and len(operations) == 1
                    await restore()
                # The web routes still fail the same way they always have: the
                # pipeline's error escapes as a plain 500.
                async with httpx.AsyncClient(
                        transport=httpx.ASGITransport(app=state.app, raise_app_exceptions=False),
                        base_url="https://test", cookies=state.client.cookies) as web:
                    for web_path, params in WEB_FILE_PATHS:
                        operations.clear()
                        with pytest.raises(AccountBindingChanged, match="account binding changed"):
                            await state.client.get(web_path, params=params)
                        assert len(operations) == 1
                        _assert_clean(state, operations)
                        await restore()
                        response = await web.get(web_path, params=params)
                        assert response.status_code == 500
                        assert response.text == "Internal Server Error"
                        assert len(operations) == 2
                        _assert_clean(state, operations)
                        await restore()
    asyncio.run(run())


def test_busy_and_resource_stops_are_busy_envelopes_and_the_web_keeps_its_own(tmp_path, monkeypatch):
    original = AdmissionManager.operation

    def admission(manager, lane, *args, **kwargs):
        if lane == "foreground":
            raise CapacityBusy("fixture admission busy")
        return original(manager, lane, *args, **kwargs)

    async def preparation_busy(*args, **kwargs):
        raise PreparationBusy("fixture preparation busy")

    async def resource_stop(*args, **kwargs):
        raise PreparationResourceError("fixture resource stop")

    async def run():
        async with application(tmp_path) as state:
            for owner, name, replacement in ((AdmissionManager, "operation", admission),
                                             (PreparationOperation, "__aenter__", preparation_busy),
                                             (PreparationOperation, "start_helper", resource_stop)):
                with monkeypatch.context() as patch:
                    patch.setattr(owner, name, replacement)
                    for path, params, web_path in FILE_PATHS:
                        response = await state.client.get(path, params=params, headers=_bearer(state))
                        assert response.status_code == 503
                        assert response.headers["retry-after"] == "1"
                        assert response.json() == _fixture("error_busy.json")
                        _assert_clean(state, [])
                        web = await state.client.get(web_path, params=params)
                        assert web.status_code == 503 and web.headers["retry-after"] == "1"
                        assert web.json() == {"ok": False, "error": "capacity_busy",
                                              "detail": "The server is busy. Please try again."}
    asyncio.run(run())


def test_a_non_ascii_pipeline_filename_is_sent_encoded(tmp_path, monkeypatch):
    original = PreparationOperation.prepared

    def prepared(operation, name, *, media_type, filename=None):
        # No pipeline names a file outside ASCII today; this stands in for one.
        return original(operation, name, media_type=media_type, filename="Fahrtenbuch ä車.xlsx")
    monkeypatch.setattr(PreparationOperation, "prepared", prepared)

    async def run():
        async with application(tmp_path) as state:
            for path, params, _ in FILE_PATHS[:2]:
                response = await state.client.get(path, params=params, headers=_bearer(state))
                assert response.status_code == 200
                assert response.headers["content-disposition"] == (
                    "attachment; filename*=utf-8''Fahrtenbuch%20%C3%A4%E8%BB%8A.xlsx")
                assert response.headers["content-length"] == str(len(response.content))
    asyncio.run(run())
