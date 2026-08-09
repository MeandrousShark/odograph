"""Tests for CSV/XLSX export."""
from __future__ import annotations

import csv
import io
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.export import HEADERS, build_export_rows, to_csv, to_range_report_xlsx, to_report_xlsx, to_xlsx
from app.expenses import build_expense_report
from app.odometer import VehicleCoverage
from app.rates import YearRate
from app.report import build_annual_report, build_range_report

TZ = ZoneInfo("America/Los_Angeles")
RATES = {2026: YearRate(0.7250)}


def _trip(**overrides) -> dict:
    base = dict(
        started_at=datetime(2026, 6, 15, 15, 0, tzinfo=timezone.utc),  # 08:00 local
        ended_at=datetime(2026, 6, 15, 15, 30, tzinfo=timezone.utc),   # 08:30 local
        distance_m=1609.344,  # exactly 1 mile
        category="business",
        purpose="Meet client",
        notes="client visit",
        has_gap=False,
        source="detected",
        start_place_name="Home",
        start_lat=47.6, start_lon=-122.3,
        end_place_name=None,
        end_lat=47.7, end_lon=-122.2,
        vehicle_name=None,
    )
    base.update(overrides)
    # display_distance_m mirrors TRIP_COLUMNS' COALESCE(distance_snapped_m,
    # distance_m); default it to distance_m unless a test overrides it
    # explicitly to simulate a snapped distance differing from the raw one.
    base.setdefault("display_distance_m", base["distance_m"])
    return base


def test_deduction_only_on_business():
    business = build_export_rows([_trip(category="business")], RATES, TZ)[0]
    personal = build_export_rows([_trip(category="personal")], RATES, TZ)[0]
    unclassified = build_export_rows([_trip(category="unclassified")], RATES, TZ)[0]
    ded_idx = HEADERS.index("Deduction ($)")
    assert business[ded_idx] == round(0.725, 2)  # 1 mi at $0.725/mi
    assert personal[ded_idx] == ""
    assert unclassified[ded_idx] == ""


def test_deduction_math():
    row = build_export_rows([_trip(category="business", distance_m=1609.344)], RATES, TZ)[0]
    ded_idx = HEADERS.index("Deduction ($)")
    assert row[ded_idx] == round(0.725, 2)


def test_deduction_uses_midyear_second_half_rate():
    # A year with a mid-year change: a July trip must price at the H2 rate,
    # exercising that build_export_rows passes the trip's local month through.
    rates = {2026: YearRate(0.700, rate_h2_per_mi=0.750, h2_start_month=7)}
    july = _trip(
        started_at=datetime(2026, 7, 15, 15, 0, tzinfo=timezone.utc),  # Jul 15 local
        ended_at=datetime(2026, 7, 15, 15, 30, tzinfo=timezone.utc),
        distance_m=1609.344,
    )
    june = _trip(distance_m=1609.344)  # default started_at is Jun 15 local
    ded_idx = HEADERS.index("Deduction ($)")
    assert build_export_rows([july], rates, TZ)[0][ded_idx] == round(0.750, 2)
    assert build_export_rows([june], rates, TZ)[0][ded_idx] == round(0.700, 2)


def test_mi_km_conversion():
    row = build_export_rows([_trip(distance_m=1609.344 * 10)], RATES, TZ)[0]
    mi_idx, km_idx = HEADERS.index("Distance (mi)"), HEADERS.index("Distance (km)")
    assert row[mi_idx] == 10.0
    assert row[km_idx] == round(1609.344 * 10 / 1000.0, 1)


def test_local_date_edge():
    # 2026-01-01 04:00 UTC is still 2025-12-31 20:00 in Los Angeles.
    trip = _trip(
        started_at=datetime(2026, 1, 1, 4, 0, tzinfo=timezone.utc),
        ended_at=datetime(2026, 1, 1, 4, 30, tzinfo=timezone.utc),
    )
    row = build_export_rows([trip], RATES, TZ)[0]
    date_idx = HEADERS.index("Date")
    assert row[date_idx] == "2025-12-31"


def test_fallback_chain_name_then_coords_then_dash():
    named = build_export_rows([_trip(start_place_name="Home")], RATES, TZ)[0]
    coords_only = build_export_rows([_trip(start_place_name=None)], RATES, TZ)[0]
    neither = build_export_rows(
        [_trip(start_place_name=None, start_lat=None, start_lon=None)], RATES, TZ
    )[0]
    start_idx = HEADERS.index("Start location")
    assert named[start_idx] == "Home"
    assert coords_only[start_idx] == "47.6000,-122.3000"
    assert neither[start_idx] == "—"


def test_export_keeps_full_reverse_geocoded_address():
    trip = _trip(
        start_place_name=None,
        start_address="123 Main St, Seattle, WA 98101",
    )
    row = build_export_rows([trip], RATES, TZ)[0]

    assert row[HEADERS.index("Start location")] == "123 Main St, Seattle, WA 98101"


def test_uses_snapped_distance_when_it_differs_from_raw():
    # A snapped trip: display_distance_m (from TRIP_COLUMNS' COALESCE)
    # differs from the noisier raw distance_m — export must use the
    # snapped figure, matching what the trip list/summaries show.
    trip = _trip(distance_m=1700.0, display_distance_m=1609.344)  # snapped to exactly 1mi
    row = build_export_rows([trip], RATES, TZ)[0]
    mi_idx = HEADERS.index("Distance (mi)")
    assert row[mi_idx] == 1.0


def test_manual_trip_row():
    manual = _trip(
        source="manual", start_place_name=None, start_lat=None, start_lon=None,
        end_place_name=None, end_lat=None, end_lon=None,
    )
    row = build_export_rows([manual], RATES, TZ)[0]
    assert row[HEADERS.index("Source")] == "manual"
    assert row[HEADERS.index("Start location")] == "—"
    assert row[HEADERS.index("End location")] == "—"


def test_vehicle_column_present_and_populated():
    with_vehicle = build_export_rows([_trip(vehicle_name="Truck")], RATES, TZ)[0]
    without_vehicle = build_export_rows([_trip(vehicle_name=None)], RATES, TZ)[0]
    vehicle_idx = HEADERS.index("Vehicle")
    assert with_vehicle[vehicle_idx] == "Truck"
    assert without_vehicle[vehicle_idx] == ""


def test_purpose_is_distinct_from_notes():
    row = build_export_rows(
        [_trip(purpose="Site inspection", notes="Parking was difficult")], RATES, TZ
    )[0]
    assert row[HEADERS.index("Purpose")] == "Site inspection"
    assert row[HEADERS.index("Notes")] == "Parking was difficult"


def test_csv_round_trip():
    trips = [_trip(category="business"), _trip(category="personal", notes=None)]
    raw = to_csv(trips, RATES, TZ)
    reader = csv.reader(io.StringIO(raw.decode("utf-8")))
    rows = list(reader)
    assert rows[0] == list(HEADERS)
    assert len(rows) == 3  # header + 2 trips
    assert rows[1][HEADERS.index("Category")] == "business"
    assert rows[2][HEADERS.index("Notes")] == ""


def test_xlsx_opens_and_has_totals_row():
    trips = [_trip(category="business", distance_m=1609.344), _trip(category="personal")]
    raw = to_xlsx(trips, RATES, TZ)
    wb = load_workbook(io.BytesIO(raw))
    ws = wb.active
    header_row = [c.value for c in ws[1]]
    assert header_row == list(HEADERS)
    assert ws.max_row == 4  # header + 2 trips + totals
    totals_row = [c.value for c in ws[ws.max_row]]
    assert totals_row[0] == "Total"
    assert totals_row[HEADERS.index("Deduction ($)")] == round(0.725, 2)


def test_report_xlsx_trips_and_summary_deduction_totals_agree():
    # Sub-cent-fraction distances (irregular increments, not round mileage)
    # so that summing already-cent-rounded per-trip deductions drifts from
    # rounding the unrounded sum once -- the exact defect this guards
    # against. 60 trips is comfortably past the >=50 the drift needs to show
    # up reliably.
    trips = [
        _trip(
            category="business",
            distance_m=(10 + i * 0.13237) * 1609.344,
            started_at=datetime(2026, 6, 15, 15, i % 59, tzinfo=timezone.utc),
            ended_at=datetime(2026, 6, 15, 15, (i % 59) + 1, tzinfo=timezone.utc),
        )
        for i in range(60)
    ]
    report = build_annual_report(trips, RATES, TZ, 2026)
    raw = to_report_xlsx(report, trips, RATES, TZ)
    wb = load_workbook(io.BytesIO(raw))

    summary_total = float(wb["Summary"]["B8"].value.lstrip("$").replace(",", ""))
    trips_ws = wb["Trips"]
    trips_total = trips_ws.cell(trips_ws.max_row, HEADERS.index("Deduction ($)") + 1).value

    expected = round(report.total_deduction, 2)
    assert summary_total == expected
    assert trips_total == expected


def test_report_xlsx_has_summary_and_trips_sheets():
    # started_at is 15:00 UTC == 08:00 local (TZ), so this trip's local month
    # is June regardless of any UTC/local boundary subtlety — not the case
    # under test here, just keeping the fixture unsurprising.
    trips = [_trip(category="business", distance_m=1609.344)]
    report = build_annual_report(trips, RATES, TZ, 2026)
    raw = to_report_xlsx(report, trips, RATES, TZ)
    wb = load_workbook(io.BytesIO(raw))

    assert wb.sheetnames == ["Summary", "Trips"]

    trips_ws = wb["Trips"]
    assert [c.value for c in trips_ws[1]] == list(HEADERS)
    assert trips_ws.cell(2, HEADERS.index("Purpose") + 1).value == "Meet client"

    summary_ws = wb["Summary"]
    assert summary_ws["A1"].value == "Annual Mileage Report — 2026"
    assert summary_ws["A3"].value == "Business miles"
    assert summary_ws["B3"].value == 1.0
    assert summary_ws["A8"].value == "Total deduction"
    assert summary_ws["B8"].value == "$0.72"  # 1mi @ $0.725/mi, float-rounded
    assert summary_ws["A11"].value == "Month"  # month table header
    assert summary_ws["A12"].value == "Jun"


def test_report_xlsx_shows_caveats_when_present():
    trips = [_trip(category="business", distance_m=1609.344, has_gap=True)]
    report = build_annual_report(trips, RATES, TZ, 2026)
    raw = to_report_xlsx(report, trips, RATES, TZ)
    wb = load_workbook(io.BytesIO(raw))
    summary_text = " ".join(
        str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value
    )
    assert "recording gap" in summary_text


def test_report_xlsx_warns_about_business_trip_without_purpose():
    trips = [_trip(category="business", purpose="")]
    report = build_annual_report(trips, RATES, TZ, 2026)
    raw = to_report_xlsx(report, trips, RATES, TZ)
    wb = load_workbook(io.BytesIO(raw))
    summary_text = " ".join(
        str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value
    )
    assert "1 business trip(s) have no purpose recorded" in summary_text


def test_report_xlsx_summary_has_by_vehicle_block():
    trips = [
        _trip(category="business", distance_m=1609.344, vehicle_name="Truck"),
        _trip(category="business", distance_m=1609.344 * 2, vehicle_name="Sedan"),
    ]
    report = build_annual_report(trips, RATES, TZ, 2026)
    raw = to_report_xlsx(report, trips, RATES, TZ)
    wb = load_workbook(io.BytesIO(raw))
    summary_rows = [
        [c.value for c in row] for row in wb["Summary"].iter_rows()
    ]
    flat = [str(v) for row in summary_rows for v in row if v is not None]
    assert "By vehicle" in flat
    assert "Truck" in flat
    assert "Sedan" in flat


def test_report_xlsx_no_caveats_block_when_clean():
    trips = [_trip(category="business", distance_m=1609.344)]
    report = build_annual_report(trips, RATES, TZ, 2026)
    raw = to_report_xlsx(report, trips, RATES, TZ)
    wb = load_workbook(io.BytesIO(raw))
    summary_text = " ".join(
        str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value
    )
    assert "Caveats" not in summary_text


def test_report_xlsx_summary_has_odometer_coverage_block_when_present():
    trips = [_trip(category="business", distance_m=1609.344)]
    report = build_annual_report(trips, RATES, TZ, 2026)
    coverage = [
        VehicleCoverage(
            vehicle_name="Truck",
            span_start=datetime(2026, 3, 1, tzinfo=TZ),
            span_end=datetime(2026, 9, 1, tzinfo=TZ),
            coverage=0.8,
            gap_m=20 * 1609.344,
            fully_bracketed=False,
        )
    ]
    raw = to_report_xlsx(report, trips, RATES, TZ, coverage)
    wb = load_workbook(io.BytesIO(raw))
    summary_text = " ".join(
        str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value
    )
    assert "Odometer reconciliation" in summary_text
    assert "Truck" in summary_text
    assert "80.0%" in summary_text


def test_report_xlsx_omits_odometer_coverage_block_when_none():
    trips = [_trip(category="business", distance_m=1609.344)]
    report = build_annual_report(trips, RATES, TZ, 2026)
    raw = to_report_xlsx(report, trips, RATES, TZ)
    wb = load_workbook(io.BytesIO(raw))
    summary_text = " ".join(
        str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value
    )
    assert "Odometer reconciliation" not in summary_text


def test_report_xlsx_has_matching_expense_comparison_and_ledger_sheet():
    from datetime import date
    from decimal import Decimal

    trips = [_trip(category="business", distance_m=1609.344, vehicle_id=7, vehicle_name="Truck")]
    expense_rows = [{
        "vehicle_id": 7,
        "vehicle_name": "Truck",
        "incurred_on": date(2026, 6, 1),
        "category": "fuel",
        "amount": Decimal("100.00"),
        "treatment": "business_use_allocated",
        "notes": "receipt 1",
    }]
    report = build_annual_report(trips, RATES, TZ, 2026)
    expense_report = build_expense_report(2026, trips, expense_rows, [], RATES, TZ)
    raw = to_report_xlsx(report, trips, RATES, TZ, None, expense_report, expense_rows)
    wb = load_workbook(io.BytesIO(raw))
    assert wb.sheetnames == ["Summary", "Trips", "Expenses"]
    summary_text = " ".join(
        str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value
    )
    assert "Standard vs. actual expense estimate" in summary_text
    assert "Provisional" in summary_text
    assert "missed driving" in summary_text
    assert [cell.value for cell in wb["Expenses"][1]] == [
        "Date", "Vehicle", "Category", "Amount ($)", "Tax treatment", "Notes",
    ]
    assert wb["Expenses"]["D2"].value == 100
    assert wb["Expenses"]["F2"].value == "receipt 1"


def test_report_xlsx_marks_inconsistent_odometer_span_ignored():
    from datetime import date
    from decimal import Decimal

    trips = [
        _trip(
            category="business", distance_m=1609.344 * 80,
            vehicle_id=7, vehicle_name="Truck",
        ),
        _trip(
            category="personal", distance_m=1609.344 * 20,
            vehicle_id=7, vehicle_name="Truck",
        ),
    ]
    expenses = [{
        "vehicle_id": 7, "vehicle_name": "Truck", "incurred_on": date(2026, 6, 1),
        "category": "fuel", "amount": Decimal("1000.00"),
        "treatment": "business_use_allocated", "notes": None,
    }]
    readings = [
        {
            "vehicle_id": 7, "vehicle_name": "Truck",
            "recorded_at": datetime(2026, 1, 1, tzinfo=TZ),
            "odometer_m": 1000 * 1609.344,
        },
        {
            "vehicle_id": 7, "vehicle_name": "Truck",
            "recorded_at": datetime(2027, 1, 1, tzinfo=TZ),
            "odometer_m": 1050 * 1609.344,
        },
    ]
    report = build_annual_report(trips, RATES, TZ, 2026)
    expense_report = build_expense_report(2026, trips, expenses, readings, RATES, TZ)
    raw = to_report_xlsx(report, trips, RATES, TZ, None, expense_report, expenses)
    wb = load_workbook(io.BytesIO(raw))
    summary_text = " ".join(
        str(cell.value) for row in wb["Summary"].iter_rows() for cell in row if cell.value
    )
    assert "Provisional (odometer ignored)" in summary_text
    assert "smaller than recorded business miles" in summary_text
    assert "80.0%" in summary_text


def test_range_report_xlsx_titles_summary_with_range_label():
    from datetime import date

    trips = [_trip(category="business", distance_m=1609.344)]
    report = build_range_report(trips, RATES, TZ, date(2026, 4, 1), date(2026, 6, 30))
    raw = to_range_report_xlsx(report, trips, RATES, TZ)
    wb = load_workbook(io.BytesIO(raw))

    assert wb.sheetnames == ["Summary", "Trips"]
    summary_ws = wb["Summary"]
    assert summary_ws["A1"].value == "Mileage Report — 2026 Q2"
    assert summary_ws["A3"].value == "Business miles"
    assert summary_ws["B3"].value == 1.0


def test_range_report_xlsx_trips_sheet_filters_out_of_range_trips():
    from datetime import date

    in_range = _trip(category="business", distance_m=1609.344, purpose="In range")
    out_of_range = _trip(
        category="business", distance_m=1609.344, purpose="Out of range",
        started_at=datetime(2026, 9, 15, 15, 0, tzinfo=timezone.utc),  # Sep local
        ended_at=datetime(2026, 9, 15, 15, 30, tzinfo=timezone.utc),
    )
    trips = [in_range, out_of_range]
    report = build_range_report(trips, RATES, TZ, date(2026, 4, 1), date(2026, 6, 30))
    raw = to_range_report_xlsx(report, trips, RATES, TZ)
    wb = load_workbook(io.BytesIO(raw))

    trips_ws = wb["Trips"]
    purposes = [
        trips_ws.cell(row, HEADERS.index("Purpose") + 1).value
        for row in range(2, trips_ws.max_row + 1)
    ]
    assert "In range" in purposes
    assert "Out of range" not in purposes


def test_range_report_xlsx_no_odometer_or_expense_sections():
    from datetime import date

    trips = [_trip(category="business", distance_m=1609.344, has_gap=True)]
    report = build_range_report(trips, RATES, TZ, date(2026, 4, 1), date(2026, 6, 30))
    raw = to_range_report_xlsx(report, trips, RATES, TZ)
    wb = load_workbook(io.BytesIO(raw))
    summary_text = " ".join(
        str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value
    )
    assert "Odometer reconciliation" not in summary_text
    assert "Standard vs. actual" not in summary_text
    # Caveats (mileage-relevant ones) still render — only the annual-only
    # sections are scoped out.
    assert "recording gap" in summary_text


def test_expenses_sheet_total_accumulates_large_values_with_decimal_precision():
    from datetime import date
    from decimal import Decimal

    amount = Decimal("9999999999.99")
    expenses = [{
        "vehicle_id": 7,
        "vehicle_name": "Truck",
        "incurred_on": date(2026, 1, 1),
        "category": "fuel",
        "amount": amount,
        "treatment": "business_use_allocated",
        "notes": None,
    } for _ in range(249)]
    expense_report = build_expense_report(2026, [], expenses, [], RATES, TZ)
    report = build_annual_report([], RATES, TZ, 2026)
    raw = to_report_xlsx(report, [], RATES, TZ, None, expense_report, expenses)
    ws = load_workbook(io.BytesIO(raw))["Expenses"]
    assert Decimal(str(ws.cell(ws.max_row, 4).value)) == Decimal("2489999999997.51")
